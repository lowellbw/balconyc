"""Rebuild the calculator's price and emissions data for US addresses.

Outside New York City the calculator values a kWh at the state's average
residential price and counts CO2 at the rate of the address's grid region.
Both come from federal spreadsheets that need no API key:

  EIA-861M sales_revenue.xlsx   monthly residential revenue and sales by state
                                (the data behind Electric Power Monthly 5.4-5.6)
  EIA Table 5A (annual)         average monthly residential use by state
  EPA Power Profiler v14.3      eGRID2023 subregion rates, and ZIP -> subregion

    python3 tools/refresh_region_data.py              # download, then build
    python3 tools/refresh_region_data.py --offline    # build from data/cache/

Writes data/state-energy.json and data/grid-emissions.json. Needs openpyxl
(pip install openpyxl). Downloads land in data/cache/, which is not deployed.

Re-run when EIA publishes a month (around the 24th) or EPA releases a new
eGRID. The price is a TRAILING-12-MONTH, sales-weighted average:
sum(revenue) / sum(sales) over the latest twelve months, so a seasonal
swing does not move it.
"""
import collections, datetime, json, pathlib, sys, urllib.request

ROOT = pathlib.Path(__file__).resolve().parents[1]
CACHE = ROOT / "data" / "cache"

SOURCES = {
    "eia861m": ("https://www.eia.gov/electricity/data/eia861m/xls/sales_revenue.xlsx",
                "sales_revenue.xlsx"),
    "table5a": ("https://www.eia.gov/electricity/sales_revenue_price/xls/table_5A.xlsx",
                "table_5A.xlsx"),
    "profiler": ("https://www.epa.gov/system/files/documents/2025-06/power_profiler_zipcode_tool_v14.3.xlsx",
                 "power_profiler_zipcode_tool_v14.3.xlsx"),
}

STATES = {
    "AL": "Alabama", "AK": "Alaska", "AZ": "Arizona", "AR": "Arkansas", "CA": "California",
    "CO": "Colorado", "CT": "Connecticut", "DE": "Delaware", "DC": "District of Columbia",
    "FL": "Florida", "GA": "Georgia", "HI": "Hawaii", "ID": "Idaho", "IL": "Illinois",
    "IN": "Indiana", "IA": "Iowa", "KS": "Kansas", "KY": "Kentucky", "LA": "Louisiana",
    "ME": "Maine", "MD": "Maryland", "MA": "Massachusetts", "MI": "Michigan", "MN": "Minnesota",
    "MS": "Mississippi", "MO": "Missouri", "MT": "Montana", "NE": "Nebraska", "NV": "Nevada",
    "NH": "New Hampshire", "NJ": "New Jersey", "NM": "New Mexico", "NY": "New York",
    "NC": "North Carolina", "ND": "North Dakota", "OH": "Ohio", "OK": "Oklahoma", "OR": "Oregon",
    "PA": "Pennsylvania", "RI": "Rhode Island", "SC": "South Carolina", "SD": "South Dakota",
    "TN": "Tennessee", "TX": "Texas", "UT": "Utah", "VT": "Vermont", "VA": "Virginia",
    "WA": "Washington", "WV": "West Virginia", "WI": "Wisconsin", "WY": "Wyoming",
}
MONTHS = ["January", "February", "March", "April", "May", "June", "July",
          "August", "September", "October", "November", "December"]
# ZIP prefixes where one three-digit prefix spans islands or grids with very
# different emissions (Oahu vs the neighbour islands; Alaska's two systems).
ZIP5_STATES_PREFIXES = ("967", "968", "995", "996", "997", "998", "999")


def fetch(key, offline):
    url, name = SOURCES[key]
    path = CACHE / name
    if not offline:
        CACHE.mkdir(parents=True, exist_ok=True)
        req = urllib.request.Request(url, headers={"User-Agent": "balco.nyc data refresh"})
        with urllib.request.urlopen(req, timeout=120) as r:
            path.write_bytes(r.read())
    if not path.exists():
        sys.exit(f"missing {path}; run without --offline to download it")
    return path


def workbook(path):
    import openpyxl
    return openpyxl.load_workbook(path, read_only=True, data_only=True)


def state_prices(path):
    ws = workbook(path)["Monthly-States"]
    rows = collections.defaultdict(dict)       # state -> {(y, m): (revenue, sales, status)}
    for r in ws.iter_rows(min_row=4, values_only=True):
        year, month, st, status, revenue, sales = r[:6]
        if st in STATES and isinstance(revenue, (int, float)) and isinstance(sales, (int, float)):
            rows[st][(int(year), int(month))] = (revenue, sales, status)
    latest = max(max(v) for v in rows.values())
    y, m = latest
    window = []
    for _ in range(12):
        window.append((y, m))
        y, m = (y - 1, 12) if m == 1 else (y, m - 1)
    window.reverse()

    out, statuses = {}, set()
    us_rev = us_sales = 0.0
    for st in STATES:
        missing = [w for w in window if w not in rows[st]]
        if missing:
            sys.exit(f"{st} lacks months {missing} in the EIA-861M file")
        rev = sum(rows[st][w][0] for w in window)          # thousand dollars
        sales = sum(rows[st][w][1] for w in window)        # MWh
        statuses.update(rows[st][w][2] for w in window)
        us_rev, us_sales = us_rev + rev, us_sales + sales
        out[st] = {"cents": round(rev / sales * 100, 2)}
    (y0, m0), (y1, m1) = window[0], window[-1]
    return out, {
        "period_label": f"12 months to {MONTHS[m1 - 1]} {y1}",
        "period_start": f"{y0}-{m0:02d}",
        "period_end": f"{y1}-{m1:02d}",
        "data_status": ", ".join(sorted(s for s in statuses if s)),
        "us_cents": round(us_rev / us_sales * 100, 2),
    }


def state_usage(path):
    ws = workbook(path)["Table 5A"]
    title = next(ws.iter_rows(max_row=1, values_only=True))[0] or ""
    by_name = {v: k for k, v in STATES.items()}
    out = {}
    for r in ws.iter_rows(min_row=4, values_only=True):
        name = (r[0] or "").strip()
        if name in by_name and isinstance(r[2], (int, float)):
            out[by_name[name]] = round(r[2], 1)
    missing = sorted(set(STATES) - set(out))
    if missing:
        sys.exit(f"Table 5A lacks {missing}")
    year = next((w for w in title.split() if w.isdigit()), None)
    return out, year


def grid(path):
    wb = workbook(path)
    subregions = {}
    for r in wb["Subregion Rates (lbs-MWh)"].iter_rows(min_row=5, values_only=True):
        name, code, co2 = r[0], r[1], r[2]
        if code and isinstance(co2, (int, float)):
            subregions[code] = {"name": name.strip(), "co2_lb_per_mwh": round(co2, 3)}

    counts = collections.defaultdict(collections.Counter)
    first = collections.defaultdict(collections.Counter)
    zip5 = {}
    for r in wb["Zip-subregion"].iter_rows(min_row=2, values_only=True):
        z = str(r[0] or "").strip().zfill(5)
        if not z.isdigit() or z < "00200":       # 00001-00199 are not USPS ZIPs
            continue
        listed = [s for s in r[1:4] if s]
        if not listed:
            continue
        for s in listed:
            counts[z[:3]][s] += 1
        first[z[:3]][listed[0]] += 1
        if z.startswith(ZIP5_STATES_PREFIXES):
            zip5[z] = listed[0]
    zip3 = {}
    for p, c in counts.items():
        # most ZIP codes; ties go to the subregion listed first more often
        zip3[p] = sorted(c, key=lambda s: (-c[s], -first[p][s], s))[0]
    return subregions, dict(sorted(zip3.items())), dict(sorted(zip5.items()))


def main(argv):
    offline = "--offline" in argv
    today = datetime.date.today().isoformat()

    prices, meta = state_prices(fetch("eia861m", offline))
    usage, usage_year = state_usage(fetch("table5a", offline))
    for st in STATES:
        prices[st]["avg_monthly_kwh"] = usage[st]
    energy = {
        "source": "U.S. Energy Information Administration, Form EIA-861M monthly sales and revenue by state "
                  "(the data behind Electric Power Monthly Tables 5.4-5.6)",
        "source_url": SOURCES["eia861m"][0],
        "method": "Residential price = sum(revenue) / sum(sales) over the latest 12 months, per state. "
                  "An average price, so it includes fixed monthly charges.",
        **meta,
        "usage_source": f"EIA Table 5A, {usage_year} average monthly residential consumption",
        "usage_source_url": SOURCES["table5a"][0],
        "retrieved": today,
        "states": prices,
    }

    subregions, zip3, zip5 = grid(fetch("profiler", offline))
    emissions = {
        "source": "EPA eGRID2023 (revision 2, June 2025) subregion total output CO2 emission rates, "
                  "and EPA Power Profiler ZIP code tool v14.3",
        "source_url": SOURCES["profiler"][0],
        "method": "zip3 = the subregion most ZIP codes in that three-digit prefix belong to. "
                  "zip5 overrides it for Hawaii and Alaska, where one prefix spans separate grids.",
        "retrieved": today,
        "subregions": subregions,
        "zip3": zip3,
        "zip5": zip5,
    }

    for name, obj in (("state-energy.json", energy), ("grid-emissions.json", emissions)):
        (ROOT / "data" / name).write_text(json.dumps(obj, indent=1, ensure_ascii=False) + "\n")
    print(f"state-energy.json: {len(prices)} states, {meta['period_label']}, US {meta['us_cents']}c")
    print(f"grid-emissions.json: {len(subregions)} subregions, {len(zip3)} ZIP3, {len(zip5)} ZIP5 overrides")


if __name__ == "__main__":
    main(sys.argv[1:])
